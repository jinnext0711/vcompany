"""テーブル図解。"""

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from src.output.helpers import _c, _style_cell, _safe, G, FS_SUB, WHITE


def render_table(slide, rect, data, design):
    """テーブルをゾーン内に描画する。"""
    headers = data.get("headers", [])
    rows = data.get("rows", [])
    if not headers:
        return

    n_cols = len(headers)
    # ゾーンに収まる行数を計算
    min_row_h = 0.38  # 11pt テキスト + パディング
    header_h = 0.42
    avail_h = _safe(rect.height - G * 2 - header_h)
    max_rows = max(int(avail_h / min_row_h), 1)
    rows = rows[:max_rows]
    n_rows = len(rows) + 1

    # テーブル高さをゾーンに収める
    table_h = min(header_h + min_row_h * len(rows), rect.height - G * 2)
    table_h = _safe(table_h, 0.5)

    shape = slide.shapes.add_table(
        n_rows, n_cols,
        Inches(rect.left), Inches(rect.top + G),
        Inches(_safe(rect.width)), Inches(table_h),
    )
    table = shape.table

    for ci, h in enumerate(headers):
        cell = table.cell(0, ci)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = _c(design, "primary")
        _style_cell(cell, 12, True, WHITE)

    for ri, row in enumerate(rows):
        for ci, val in enumerate(row[:n_cols]):
            cell = table.cell(ri + 1, ci)
            cell.text = val
            if ci == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = _c(design, "light_bg")
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = (
                    _c(design, "light_bg") if ri % 2 == 1
                    else RGBColor(255, 255, 255)
                )
            _style_cell(cell, FS_SUB, ci == 0, _c(design, "text"))
